from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

root=Path(r'C:\Users\intel\Documents\Playground\geodiff_gan')
out=root/'reports'; out.mkdir(exist_ok=True)
path=out/'GeoDiff_GAN_Progress_Report_June_2026.docx'
doc=Document(); sec=doc.sections[0]
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
sec.header_distance=sec.footer_distance=Inches(.492)
styles=doc.styles
normal=styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(11); normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.1
for name,size,color,before,after in [('Title',24,'0B2545',0,6),('Heading 1',16,'2E74B5',16,8),('Heading 2',13,'2E74B5',12,6),('Heading 3',12,'1F4D78',8,4)]:
 s=styles[name]; s.font.name='Calibri'; s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(color); s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after)
# header/footer
h=sec.header.paragraphs[0]; h.text='GeoDiff-GAN Research Progress'; h.runs[0].font.size=Pt(9); h.runs[0].font.color.rgb=RGBColor(100,100,100)
f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=f.add_run('Progress report | June 2026'); r.font.size=Pt(9); r.font.color.rgb=RGBColor(100,100,100)
# title
p=doc.add_paragraph(style='Title'); p.add_run('GeoDiff-GAN: Research Progress Report')
p=doc.add_paragraph(); r=p.add_run('Prompt-Guided Diffusion-GAN Super-Resolution for Sentinel-2 Satellite Imagery'); r.bold=True; r.font.size=Pt(14); r.font.color.rgb=RGBColor.from_string('1F4D78')
for label,value in [('Prepared for','Research Supervisor'),('Prepared by','Shashank'),('Reporting date','23 June 2026'),('Research status','Architecture implemented; controlled training and comparative evaluation in progress')]:
 p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(2); a=p.add_run(label+': '); a.bold=True; p.add_run(value)

def add_bullets(items):
 for x in items: doc.add_paragraph(x, style='List Bullet')
def shade(cell,fill):
 tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def table(headers,rows,widths=None):
 t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'; t.autofit=False
 for i,hv in enumerate(headers):
  c=t.rows[0].cells[i]; c.text=hv; shade(c,'F2F4F7'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  for rr in c.paragraphs[0].runs: rr.bold=True
 for row in rows:
  cells=t.add_row().cells
  for i,v in enumerate(row): cells[i].text=str(v); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
 if widths:
  for row in t.rows:
   for i,w in enumerate(widths): row.cells[i].width=Inches(w)
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
 return t

doc.add_heading('1. Executive Summary',1)
doc.add_paragraph('This work investigates 4x single-image super-resolution for Sentinel-2 L2A RGB imagery, using synthetic 40 m observations and native 10 m targets. The proposed GeoDiff-GAN separates conservative reconstruction from uncertain detail synthesis. A deterministic SwinIR branch establishes radiometry and geometry, latent diffusion models plausible residual detail, a dual-policy GeoMapper controls evidence confidence and edit permission, and a residual GAN decoder reconstructs high-frequency content. Sensor back-projection finally enforces consistency with the observed LR image.')
add_bullets(['Implemented an end-to-end multi-stage research codebase with data preparation, training, evaluation, uncertainty, diagnostics, and resumable notebooks.', 'Developed small, medium, and large variants, plus an improved small model using resize-convolution to reduce periodic artifacts.', 'Completed six-tile diagnostic training and expanded the controlled experiment to twelve products with validation-based checkpointing and early stopping.', 'Built a common benchmark harness for recent open-source SR models and a saved-checkpoint evaluation notebook covering PSNR, SSIM, LPIPS, DISTS, edge F1, and LR consistency.'])

doc.add_heading('2. Research Problem and Scope',1)
doc.add_paragraph('The primary experiment maps a 128 x 128 RGB LR patch to a 512 x 512 RGB HR patch. Because 4x degradation is non-invertible, the system cannot guarantee unique recovery of missing sub-pixel structure. The defensible objective is evidence-constrained reconstruction with explicit uncertainty, not unrestricted hallucination.')
table(['Item','Current scope'],[['Sensor/product','Sentinel-2 L2A surface reflectance'],['Bands','B4/B3/B2 RGB'],['Task','Synthetic 40 m to native 10 m, 4x SR'],['HR patch','3 x 512 x 512'],['LR patch','3 x 128 x 128'],['Modes','SR reconstruction and explicitly synthetic prompt edit'],['Split principle','Product/tile-aware separation; spatial-block split used for controlled DGX comparison']], [1.6,4.9])

doc.add_heading('3. Proposed Architecture',1)
img=root/'docs/images/geodiff-dual-policy-architecture.png'
if img.exists():
 doc.add_picture(str(img), width=Inches(6.4)); p=doc.paragraphs[-1]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 cap=doc.add_paragraph('Figure 1. GeoDiff-GAN dual-policy architecture and evidence-controlled residual path.'); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER; cap.runs[0].italic=True

table(['Module','Input -> output','Research role'],[
['Deterministic base','3x128x128 -> 3x512x512','Conservative geometry, reflectance, and low-frequency reconstruction.'],
['Residual VAE','3x512x512 residual -> 4x64x64 latent','Compresses the target residual distribution for latent diffusion.'],
['LR encoder','3x128x128 -> multi-scale features','Anchors latent and decoder features to observed spatial evidence.'],
['Conditional diffusion','4x64x64 noisy latent -> denoised latent','Models ambiguous, stochastic residual detail conditioned on LR and degradation.'],
['Dual-policy GeoMapper','latent + LR + context -> content, styles, confidence, permission','Separates evidence-supported reconstruction from prompt-controlled synthetic editing.'],
['Residual decoder','64-grid content -> 3x512x512 residual','Generates detail residuals with LR skip connections and layer-wise modulation.'],
['Back-projection','HR estimate + LR -> corrected HR','Reduces sensor re-degradation error while preserving the residual formulation.'],
['Discriminators','HR/LR pairs during training only','Patch realism and Haar high-frequency supervision; absent at inference.']], [1.25,1.65,3.6])

doc.add_heading('4. Novel Research Contributions Under Evaluation',1)
add_bullets(['Evidence-controlled residual generation: stochastic detail is added to a deterministic base rather than generating an unrestricted HR image.', 'Dual spatial policies: evidence confidence controls SR detail, while edit permission independently controls prompt-driven synthetic changes.', 'Uncertainty-aware abstention: multi-sample variance reduces trust in unstable generated detail and blends toward the deterministic base.', 'Sensor-consistency closure: iterative back-projection uses the degradation model to constrain the final output to the LR observation.', 'Dual-frequency adversarial training: conditional multi-scale PatchGAN is combined with a Haar-wavelet discriminator.', 'Explicit SR/edit semantics: reconstruction outputs and synthetic prompt edits use different residual policies and metadata claims.'])
doc.add_paragraph('These are proposed contributions, not established novelty claims. Each requires controlled ablation against deterministic, diffusion-only, GAN-only, no-gate, no-projection, and alternative-decoder baselines.')

doc.add_heading('5. Model Variants and Capacity',1)
table(['Variant','Core parameters','Purpose'],[['XS smoke','0.765 M','Pipeline and shape testing only'],['Small','12.140 M','Primary constrained experiments and debugging'],['Small improved','12.054 M','Resize-convolution decoder and stronger calibration/detail supervision'],['Medium','21.127 M','Preferred capacity-scaling experiment'],['Large','81.856 M','Full-capacity final experiment after data and decoder validation']], [1.5,1.5,3.5])

doc.add_heading('6. Data and Training Pipeline',1)
add_bullets(['Complete Sentinel-2 SAFE products are read window-by-window; accepted patches are stored as compressed NPZ files.', 'SCL filtering removes cloud, shadow, invalid, saturated, and edge/corner background patches; rejected data are quarantined for audit.', 'Training LR is generated online with randomized MTF/PSF blur, area downsampling, mild Poisson-Gaussian noise, quantization, and degradation parameters.', 'Training follows base, residual VAE/decoder, latent diffusion, joint diffusion-GAN, and optional prompt-edit stages.', 'Mixed precision, gradient accumulation, resumable checkpoints, compact progress reporting, validation-selected best checkpoints, and early stopping are implemented.', 'Offline caption generation supports brief, descriptive, and analytical captions; a grounded multispectral workflow using B08/B11/SCL evidence was designed to reduce RGB-only caption hallucination.'])

doc.add_heading('7. Experimental Progress and Findings',1)
doc.add_heading('Six-product diagnostic',2)
table(['Diagnostic','Observed value','Interpretation'],[['Base-to-target L1','0.04736','Deterministic base remained inaccurate on the selected validation patch.'],['Output-to-target L1','0.02650','Final system improved patch-level reconstruction by 44%.'],['LR error before/after projection','0.02032 / 0.00329','Back-projection reduced inconsistency by 83.8%.'],['Evidence confidence / abstention','0.166 / 0.834','The model suppressed most proposed stochastic detail.'],['Output/target edge magnitude','33.3%','High-frequency reconstruction was substantially too weak.'],['Wavelet amplitude retained','30.7-37.8%','Directional and diagonal detail remained under-reconstructed.'],['Numerical stability','No NaN/Inf','Pipeline and diffusion sampling were stable.']], [1.6,1.35,3.55])
doc.add_paragraph('Conclusion: the model was not numerically failing. It improved the deterministic base and preserved LR evidence, but remained overly conservative and failed to reconstruct enough target high-frequency structure.')
img=root/'learning/images/small_six_tile_diagnostics/01_overview.png'
if img.exists():
 doc.add_picture(str(img), width=Inches(6.35)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
 cap=doc.add_paragraph('Figure 2. Representative six-product validation diagnostic: LR, base, residual, output, target, and consistency maps.'); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER; cap.runs[0].italic=True

doc.add_heading('Twelve-product improved experiment',2)
add_bullets(['Replaced residual-decoder PixelShuffle stages with bilinear resize-convolution to reduce lattice/checkerboard artifacts.', 'Added spatial confidence-selectivity supervision to discourage nearly uniform evidence maps.', 'Strengthened gradient and wavelet supervision while retaining a low adversarial weight.', 'Added 0/1/3-step projection ablations, early stopping, and best-plus-latest checkpoint retention.', 'Observed joint validation peaked early in one run, indicating overfitting or an overly aggressive joint phase; epoch 1 was retained as the best checkpoint.'])

doc.add_heading('8. Current Limitations and Risks',1)
add_bullets(['Insufficient geographic diversity in early experiments limits generalization claims.', 'High-frequency detail remains weaker than the HR target even when LR consistency is strong.', 'Back-projection can repair LR consistency but cannot prove that inferred HR edges are correct.', 'Evidence confidence may become globally conservative instead of spatially selective.', 'Prompt captions generated from RGB alone can misclassify dark land as water; multispectral grounding is required.', 'Synthetic 40 m degradation does not prove performance on native real 40 m imagery.', 'Large capacity on small datasets increases overfitting risk and does not automatically fix decoder artifacts.'])

doc.add_heading('9. Benchmarking and Evaluation',1)
doc.add_paragraph('A common evaluation framework has been prepared for GeoDiff-GAN and recent open-source SR architectures, including SwinIR, HAT, SRFormer, DAT, OmniSR, TTST, MFG-HMoE, and optional FreMamba. Saved checkpoints can now be evaluated without retraining using identical deterministic LR inputs.')
table(['Metric','Meaning','Direction'],[['PSNR','Pixel-level fidelity; can favor smooth outputs','Higher'],['SSIM','Local luminance, contrast, and structural similarity','Higher'],['LPIPS','Deep-feature perceptual distance','Lower'],['DISTS','Structure and texture perceptual distance','Lower'],['Edge F1','Agreement of reconstructed and target edges','Higher'],['Re-degradation L1','Consistency of SR output with observed LR','Lower'],['Uncertainty correlation','Whether stochastic variance identifies difficult regions','Positive/error-aware']], [1.35,3.8,1.35])

doc.add_heading('10. Next Work Plan',1)
for i,text in enumerate(['Run saved-checkpoint comparison on the complete held-out split with LPIPS and DISTS enabled.', 'Compare small improved and medium under identical data, updates, degradation seed, and early-stopping policy.', 'Complete decoder, evidence-gate, text, wavelet-discriminator, degradation-conditioning, and projection ablations.', 'Increase geographically diverse products while keeping validation and test tiles isolated.', 'Audit grounded captions by spectral evidence and manually review high-risk water/cloud/forest cases.', 'Report eight stochastic samples per test patch with uncertainty maps and selective-risk analysis.', 'Use the large model only after medium demonstrates reliable held-out improvement.'],1):
 doc.add_paragraph(text, style='List Number')

doc.add_heading('11. Supervisor Decisions Requested',1)
add_bullets(['Approve the medium model as the next capacity experiment before committing compute to the large model.', 'Confirm that final claims should focus on synthetic Sentinel-2 40 m to 10 m reconstruction, with real native-resolution transfer treated as future validation.', 'Confirm the minimum number and geographic distribution of tiles required for thesis-level evaluation.', 'Review whether prompt-guided edit mode should remain a secondary synthetic-visualization contribution or be excluded from the first paper.'])

doc.add_heading('12. Current Defensible Research Statement',1)
p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.25); p.paragraph_format.right_indent=Inches(.25)
r=p.add_run('GeoDiff-GAN is an evidence-controlled residual super-resolution framework that combines a deterministic transformer base, conditional latent diffusion, dual reconstruction/edit policies, adversarial high-frequency supervision, uncertainty abstention, and sensor back-projection. Current experiments show stable execution, improved reconstruction over the base, and strong LR consistency, while also revealing insufficient high-frequency recovery and the need for broader geographic validation and controlled ablation.'); r.italic=True

doc.save(path)
print(path)
