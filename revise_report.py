from pathlib import Path
import shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(r'C:\Users\intel\Documents\Playground\geodiff_gan')
src=root/'reports/GeoDiff_GAN_Progress_Report_June_2026.docx'
dst=root/'reports/GeoDiff_GAN_Progress_Report_June_2026_Revised.docx'
shutil.copy2(src,dst)
doc=Document(dst)
body=doc._element.body
for child in list(body):
    if child.tag != qn('w:sectPr'):
        body.remove(child)

def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def tbl(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,'F2F4F7'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs:r.bold=True
    for row in rows:
        cs=t.add_row().cells
        for i,v in enumerate(row): cs[i].text=str(v); cs[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph()
    return t
def bullets(items):
    for item in items: doc.add_paragraph(item,style='List Bullet')
def fig(path,caption,width=6.4):
    doc.add_picture(str(path),width=Inches(width)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph(caption); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].italic=True; p.runs[0].font.size=Pt(9)

p=doc.add_paragraph(style='Title'); p.add_run('GeoDiff-GAN: Revised Research Progress Report')
p=doc.add_paragraph(); r=p.add_run('Evidence-Constrained Diffusion-GAN Super-Resolution for Sentinel-2 Imagery'); r.bold=True; r.font.size=Pt(14); r.font.color.rgb=RGBColor.from_string('1F4D78')
for label,value in [('Prepared for','Research Supervisor'),('Prepared by','Shashank'),('Reporting date','23 June 2026'),('Status','Architecture implemented; trained variants and open-source baselines evaluated')]:
 p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(2); a=p.add_run(label+': '); a.bold=True; p.add_run(value)

doc.add_heading('1. Executive Summary',1)
doc.add_paragraph('GeoDiff-GAN addresses 4x Sentinel-2 RGB super-resolution from synthetic 40 m observations to native 10 m targets. The architecture deliberately separates reliable low-frequency reconstruction from uncertain high-frequency synthesis. A deterministic SwinIR anchor reconstructs geometry and radiometry; conditional latent diffusion proposes residual detail; the GeoMapper controls evidence confidence; a residual decoder adds detail; and sensor back-projection closes the loop with the LR observation.')
doc.add_paragraph('The current results are defensible as a strong progress result rather than a final state-of-the-art claim. On the common ten-image comparison, the medium model achieved the best LPIPS (0.2349) and edge F1 (0.559), while the improved small model was second on both measures. These results show that the proposed losses, evidence-controlled residual pathway, and resize-convolution decoder recover perceptually clearer and more spatially explicit features than the tested deterministic baselines. HAT and OmniSR retained higher PSNR/SSIM, illustrating the expected perception-distortion trade-off rather than invalidating the proposed method.')

doc.add_heading('2. Task Definition',1)
tbl(['Item','Current experiment'],[['Input','Sentinel-2 RGB, 3 x 128 x 128, synthetic 40 m'],['Target/output','RGB, 3 x 512 x 512, native 10 m'],['Scale','4x single-image super-resolution'],['Operating mode reported here','Evidence-constrained SR with empty/null prompt'],['Evaluation split','Ten held-out images in the saved-model comparison'],['Claim boundary','Synthetic-degradation reconstruction; not unique recovery of unobserved physical truth']], [1.8,4.7])

doc.add_heading('3. Proposed Inference Architecture',1)
arch=Path(r'C:\Users\intel\AppData\Local\Temp\codex-clipboard-89068419-6737-440b-8102-b9321d636d16.png')
fig(arch,'Figure 1. Evidence-SR inference graph for the practical medium model with an empty prompt.',6.5)
doc.add_paragraph('The diagram distinguishes inference modules from faded training-only supervision. During inference, discriminators and the HR/VAE target encoder are absent. The observed LR follows two paths: a deterministic SwinIR anchor and a multi-scale LR evidence pyramid. Conditional diffusion generates a latent residual proposal, while GeoMapper combines latent content, degradation condition, mode condition, and LR evidence. The dual-head decoder predicts residual detail, the SR policy retains only evidence-supported high-pass content, and sensor back-projection produces an LR-consistent HR result.')
tbl(['Module','Role in evidence preservation'],[['SwinIR base','Provides a stable HR anchor and prevents the generative branch from controlling all image content.'],['LR encoder','Carries measured spatial evidence at 128, 64, 32, and 16 grids.'],['Conditional diffusion','Models multiple plausible residual-detail solutions rather than averaging every ambiguity.'],['GeoMapper','Maps latent content to decoder content/styles and predicts evidence policy.'],['Dual-head decoder','Separates shared/detail reconstruction from edit behavior; only the SR residual is used here.'],['Residual composition','Adds high-pass evidence-weighted detail instead of unrestricted full-band generation.'],['Sensor back-projection','Reduces disagreement between the final HR image and the measured LR observation.']], [1.55,4.95])

doc.add_heading('4. Architecture and Training Improvements',1)
bullets(['PixelShuffle was removed from the improved residual decoder and replaced by resize-convolution. This directly targets phase imbalance and repeated lattice/checkerboard patterns seen in Fourier diagnostics.', 'Gradient and wavelet losses were strengthened and applied to the VAE/joint reconstruction paths. These losses make roads, field boundaries, settlement blocks, ridges, and fine directional structures more visible instead of optimizing only average pixel agreement.', 'Evidence-confidence calibration was extended with a spatial selectivity term, discouraging the previous nearly uniform low-confidence solution.', 'The adversarial weight remained low. This avoids obtaining visually sharp but unsupported textures while the decoder and confidence policy are still being calibrated.', 'Early stopping and validation-selected best checkpoints prevent the joint stage from continuing after held-out quality begins to decline.', 'Projection-step ablations were added because strong LR consistency alone can hide weak learned high-frequency reconstruction.'])
doc.add_paragraph('Qualitatively, the loss changes and removal of PixelShuffle made scene-dependent residuals more visible and reduced the earlier tendency to produce nearly identical periodic residual texture across different inputs. The output still does not exactly match the ground truth, which is expected for an ill-posed 4x inverse problem, but the reconstructed structures are more image-specific and perceptually meaningful.')

doc.add_heading('5. Quantitative Comparison',1)
fig(root/'results/table.png','Figure 2. Saved-model comparison table on ten held-out samples.',6.5)
tbl(['Method','PSNR','SSIM','Edge F1','LPIPS','LR re-degradation L1'],[['GeoDiff medium','32.000','0.849','0.559','0.2349','0.001073'],['GeoDiff small improved','32.439','0.845','0.533','0.2535','0.001142'],['HAT','33.501','0.875','0.372','0.3903','0.003231'],['OmniSR','33.465','0.874','0.368','0.3930','0.003243'],['MFG-HMoE','33.130','0.807','0.304','0.4249','0.003527'],['SRFormer','32.040','0.860','0.254','0.4921','0.009336'],['SwinIR','31.219','0.856','0.261','0.4982','0.013610']], [1.6,.75,.75,.75,.75,1.4])
fig(root/'results/colab_metric_download.png','Figure 3. Metric comparison. GeoDiff leads perceptual distance and edge reconstruction; HAT/OmniSR lead distortion metrics.',6.5)

doc.add_heading('6. Importance of Each Metric and Defense of the Results',1)
tbl(['Metric','Why it matters','How the current result should be interpreted'],[['PSNR (higher)','Measures pixelwise fidelity through MSE; important for radiometric reconstruction but rewards smooth averages.','GeoDiff is competitive but below HAT/OmniSR. This is consistent with adding more high-frequency structure instead of optimizing only smooth pixel agreement.'],['SSIM (higher)','Measures local luminance, contrast, and structural similarity.','GeoDiff remains structurally credible, although deterministic transformers score higher on this small test.'],['Edge F1 (higher)','Measures whether roads, roofs, boundaries, and linear structures occur at matching locations.','This is GeoDiff’s strongest result: medium 0.559 and small improved 0.533, substantially above HAT 0.372.'],['LPIPS (lower)','Measures perceptual distance in deep feature space and is sensitive to meaningful texture/structure.','GeoDiff medium is best at 0.2349; improved small is second at 0.2535. This supports visibly clearer, more target-like features.'],['LR re-degradation L1 (lower)','Tests whether the HR prediction returns to the observed LR under the sensor degradation model.','GeoDiff is best by a wide margin, supporting the evidence-constrained design and back-projection.'],['L1 (lower)','Direct mean absolute HR reconstruction error.','GeoDiff medium 0.0158 and small improved 0.0161 are competitive, though HAT/OmniSR remain lower.']], [1.15,2.25,3.1])
doc.add_paragraph('The combined interpretation is more important than any single ranking. HAT and OmniSR produce the strongest conventional distortion scores, while GeoDiff produces substantially better edge agreement, perceptual similarity, and sensor consistency. For satellite SR, this trade-off is relevant because a high-PSNR image may remain too smooth to reveal narrow roads, settlement boundaries, or field structure. Conversely, edge and LPIPS improvements must always be checked against LR consistency to avoid defending hallucinated sharpness. GeoDiff performs well on both axes: it sharpens meaningful features while achieving the lowest re-degradation error.')
doc.add_paragraph('The results are based on ten held-out images and therefore demonstrate promising behavior, not final statistical superiority. The next report should include the complete test split, confidence intervals or per-tile distributions, and repeated stochastic sampling.')

doc.add_heading('7. Qualitative Results',1)
fig(root/'results/colab_infer_download.png','Figure 4. Side-by-side inference outputs and absolute-error maps for saved GeoDiff and baseline models.',6.1)
doc.add_paragraph('The qualitative comparison should be read together with Figure 3. After modifying the loss balance and removing PixelShuffle, roads, texture transitions, agricultural boundaries, and settlement patterns become more visible and input-dependent. The improved residual no longer behaves like a common repeating texture template. Remaining differences from the target occur mainly in fine high-frequency details that are not uniquely determined by the 40 m observation.')

doc.add_heading('8. Benchmark Scope and Exclusions',1)
bullets(['DISTS was removed from the final comparison. The report therefore uses LPIPS as the deep perceptual metric and does not present incomplete or inconsistent DISTS values.', 'TTST is not included in the quantitative table because its public GitHub repository did not run correctly in the available environment. It would be misleading to report a failed or modified implementation as a faithful TTST baseline.', 'The evaluated open-source baselines are SwinIR, SRFormer, HAT, OmniSR, and the remote-sensing-specific MFG-HMoE implementation.', 'All reported baseline checkpoints were trained or evaluated on the same prepared Sentinel-2 task rather than copying scores from unrelated natural-image datasets.', 'The comparison is specific to the saved split, degradation model, and evaluation code; it is not a universal 2026 leaderboard.'])

doc.add_heading('9. Current Limitations',1)
bullets(['The ten-image comparison is too small for a final generalization claim.', 'Synthetic 40 m degradation may not perfectly reproduce a native sensor’s real blur, noise, atmosphere, or registration errors.', 'Back-projection enforces observation consistency but cannot prove that every generated 10 m edge is physically correct.', 'Prompt conditioning has not yet been included in this saved-model comparison.', 'Caption generation from RGB alone caused land/water confusion; future prompt experiments should use multispectral grounding and manual audit.', 'Medium and improved-small results should be repeated over complete geographically separated test tiles.'])

doc.add_heading('10. Next Work',1)
for item in ['Evaluate all saved checkpoints on the complete held-out test set with per-tile results.', 'Report PSNR, SSIM, L1, edge F1, LPIPS, and re-degradation L1 with consistent inference settings.', 'Run ablations for resize-convolution, gradient/wavelet weights, evidence calibration, and 0/1/3 back-projection steps.', 'Inspect frequency spectra to verify that PixelShuffle-related lattice peaks remain reduced.', 'Compare improved small and medium with identical update budgets; use the large model only if medium scaling improves held-out results.', 'Generate grounded captions using B08, B11, spectral indices, and SCL evidence before prompt-conditioned training.']:
 doc.add_paragraph(item,style='List Number')

doc.add_heading('11. Defensible Conclusion',1)
p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.25); p.paragraph_format.right_indent=Inches(.25)
r=p.add_run('GeoDiff-GAN currently demonstrates a meaningful perception-and-evidence advantage: the medium model achieves the best LPIPS, edge F1, and LR re-degradation consistency among the evaluated saved models, while remaining competitive in PSNR and SSIM. Removing PixelShuffle and strengthening controlled detail losses made reconstructed features more visible and scene-specific. These findings support continued evaluation of the evidence-controlled residual architecture, but final novelty and superiority claims require complete geographically separated testing and controlled ablations.'); r.italic=True

doc.save(dst)
print(dst)
